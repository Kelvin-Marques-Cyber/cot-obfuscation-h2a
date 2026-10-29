"""Gera o artigo em DOCX (editável) e PDF a partir de revisao_2026/artigo.md.

Fluxo: figuras (gerar_figuras.py) -> expansão de tabelas, figuras, valores e
referências cruzadas -> pandoc com modelo de estilos -> pós-processamento das
tabelas -> LibreOffice converte o MESMO DOCX em PDF -> verificações (páginas,
fontes incorporadas, equivalência textual DOCX/PDF) -> PNG de todas as páginas.

Uso (a partir da raiz do repositório):
    python revisao_2026/src/gerar_artigo.py
Requer pandoc 3, LibreOffice com o componente Math, poppler-utils, python-docx
e matplotlib.
"""
from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
REPO = ROOT.parent
SOURCE = ROOT / 'artigo.md'
BUILD = ROOT / 'build'
DOCS = ROOT / 'docs'
NAME = 'Artigo_CoT_Versao_Final'
sys.path.insert(0, str(Path(__file__).resolve().parent))


# --------------------------------------------------------------- formatação
def num(value, digits=2):
    text = f'{value:,.{digits}f}'
    return text.replace(',', '§').replace('.', ',').replace('§', '.')


def pct(value, digits=1):
    return num(100 * value, digits) + '%'


def frac(k, n):
    return f'{k}/{n}'


def wilson(k, n, z=1.959963984540054):
    """Intervalo de confiança de Wilson para uma proporção."""
    p = k / n
    centre = (p + z * z / (2 * n)) / (1 + z * z / n)
    half = z * ((p * (1 - p) / n + z * z / (4 * n * n)) ** 0.5) / (1 + z * z / n)
    return centre - half, centre + half


def read_csv(path):
    with Path(path).open(encoding='utf-8') as handle:
        return list(csv.DictReader(handle))


def read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


# ------------------------------------------------------------------- dados
DIAG = REPO / 'experimento_cientifico_v3' / 'diagnostico_dev' / 'resultados' / 'diagnostico_triagem_dev.json'
V4 = REPO / 'experimento_cientifico_v4'
VALIDATION = ROOT / 'results' / 'validacao_2026_10_28.json'


def values():
    """Todos os números citados no texto vêm daqui (chave -> texto formatado)."""
    out = {}
    report = read_json(ROOT / 'results' / 'report.json')
    summary = read_csv(ROOT / 'results' / 'summary.csv')
    contrasts = {r['condition']: r for r in read_csv(ROOT / 'results' / 'contrasts.csv')}
    stress = read_csv(ROOT / 'results' / 'grading_stress_summary.csv')
    out['piloto.linhas'] = str(report['n_rows'])
    out['piloto.celulas'] = str(report['n_cells'])
    out['piloto.anova_f'] = num(report['historical_anova']['F'], 2)
    out['piloto.anova_p'] = num(report['historical_anova']['p'] * 1e12, 2) + ' × 10⁻¹²'
    out['piloto.welch_t'] = num(report['historical_welch']['t'], 2)
    out['piloto.welch_p'] = num(report['historical_welch']['p'] * 1e7, 2) + ' × 10⁻⁷'
    for row in summary:
        out[f'piloto.acc.{row["level"]}.{row["condition"]}'] = num(float(row['mean']), 3)
    for condition, row in contrasts.items():
        out[f'piloto.dd.{condition}'] = num(float(row['mean']), 3)
        out[f'piloto.dd_ic.{condition}'] = f'[{num(float(row["conditional_t_low"]), 3)}; {num(float(row["conditional_t_high"]), 3)}]'
        out[f'piloto.signflip.{condition}'] = num(float(row['sign_flip_two_sided']), 2)
    necessity = report['necessity_reconstructed']['C']
    out['piloto.mcnemar_b'] = str(necessity['cot_only'])
    out['piloto.mcnemar_c'] = str(necessity['direct_only'])
    out['piloto.mcnemar_p'] = num(necessity['p_exact_mcnemar'], 3)
    out['piloto.ganho_liquido'] = num(necessity['net_gain'], 2)
    out['piloto.falsos_aceites'] = num(report['original_false_accepts'], 0)
    out['piloto.negativos'] = num(report['negative_cases'], 0)
    out['piloto.casos'] = num(report['constructed_stress_cases'], 0)
    out['piloto.taxa_aceite_indevido'] = pct(report['original_false_accepts'] / report['negative_cases'])
    out['piloto.itens_disjuntos'] = num(report['disjoint_generated_items'], 0)
    out['piloto.corr_rep'] = num(report['correlations']['score_repetition'], 3)
    out['piloto.corr_ttr'] = num(report['correlations']['score_ttr'], 3)
    del stress
    if DIAG.exists():
        diag = read_json(DIAG)
        for key, entry in diag['levels'].items():
            mode, level = key.split('/')
            base = f'v3.{mode}.{level}'
            t4 = entry['t4_fp16']
            out[base + '.doc'] = pct(entry['documented_accuracy'], 2 if mode == 'grpo' else 1)
            out[base + '.doc_k'] = frac(entry['documented_correct'], 120)
            out[base + '.t4'] = pct(t4['accuracy'])
            out[base + '.t4_k'] = frac(t4['correct'], 120)
            out[base + '.perm'] = pct(entry['permissive_upper_bound'])
            out[base + '.t4_trunc'] = str(entry['t4_truncated_estimated'])
            for name, count in entry['categories'].items():
                out[f'{base}.cat.{name}'] = str(count)
            cpu = entry.get('cpu_fp32')
            if cpu:
                out[base + '.cpu_k'] = frac(cpu['correct'], cpu['n'])
                out[base + '.cpu_n'] = str(cpu['n'])
                out[base + '.cpu_ident'] = str(cpu['identical_generation'])
                out[base + '.cpu_igual'] = str(cpu['same_correctness'])
                out[base + '.trunc'] = str(cpu['truncated'])
                out[base + '.tokens'] = num(cpu['mean_output_tokens'], 1)
            if mode == 'grpo':
                out[base + '.formato'] = str(entry['format_valid'])
                out[base + '.dica'] = str(entry['chose_hint_letter_valid_format'])
                out[base + '.marcador'] = str(entry['copied_instruction_placeholder'])
                out[base + '.letra'] = str(entry['answer_letra_literal'])
                for domain, d in entry['by_domain'].items():
                    out[f'{base}.dom.{domain}'] = frac(d['correct'], d['n'])
    cal = V4 / 'calibracao' / 'resultados'
    for name, prefix in (('sft_calibration_decision.json', 'cal.sft.r1'), ('sft_calibration_decision_r2.json', 'cal.sft.r2')):
        if (cal / name).exists():
            for key, entry in read_json(cal / name)['summary'].items():
                tag = key.replace('/', '.')
                out[f'{prefix}.{tag}'] = pct(entry['accuracy'])
                out[f'{prefix}.{tag}.k'] = frac(round(entry['accuracy'] * entry['n']), entry['n'])
    for name, prefix in (('grpo_calibration_decision.json', 'cal.grpo.r1'), ('grpo_calibration_decision_r2.json', 'cal.grpo.r2')):
        if (cal / name).exists():
            for key, levels in read_json(cal / name)['summary'].items():
                for level, entry in levels.items():
                    tag = key.replace('/', '.') + '.' + level
                    out[f'{prefix}.{tag}'] = pct(entry['accuracy'])
                    out[f'{prefix}.{tag}.k'] = frac(round(entry['accuracy'] * entry['n']), entry['n'])
                    out[f'{prefix}.{tag}.dica'] = pct(entry['chose_hint'])
                    out[f'{prefix}.{tag}.formato'] = pct(entry['format_valid'])
    short = {'HuggingFaceTB/SmolLM2-135M-Instruct': '135m', 'HuggingFaceTB/SmolLM2-360M-Instruct': '360m',
             'HuggingFaceTB/SmolLM2-1.7B-Instruct': 'smol17b', 'Qwen/Qwen2.5-0.5B-Instruct': 'reward',
             'Qwen/Qwen2.5-1.5B-Instruct': 'qwen15b', 'allenai/OLMo-2-0425-1B-Instruct': 'olmo1b',
             'ibm-granite/granite-3.3-2b-instruct': 'granite2b', 'HuggingFaceTB/SmolLM3-3B': 'smol3',
             'microsoft/Phi-3.5-mini-instruct': 'phi35'}
    for name in ('monitor_feasibility_decision.json', 'monitor_feasibility_decision_r2.json'):
        if (cal / name).exists():
            for entry in read_json(cal / name)['summary']:
                key = short[entry['model_id']]
                if entry['tpr'] is None:
                    out[f'mon.{key}.status'] = 'não avaliável na T4 em fp32'
                    continue
                out[f'mon.{key}.tpr'] = pct(entry['tpr'])
                out[f'mon.{key}.fpr'] = pct(entry['fpr'])
                for template, rate in entry.get('flag_rate_by_template', {}).items():
                    out[f'mon.{key}.{template}'] = pct(rate)
    r2a = cal / 'grpo_G1_sem_dica_r2a.jsonl'
    if r2a.exists():
        rows = [json.loads(line) for line in r2a.read_text(encoding='utf-8').splitlines() if line.strip()]
        lenient = re.compile(r'<answer>\s*([A-D])\s*</answer>\s*$')
        out['cal.r2a.n'] = str(len(rows))
        out['cal.r2a.quebra'] = str(sum('</trace>\n' in r['generation'] or '</summary>\n' in r['generation']
                                        for r in rows))
        out['cal.r2a.letra'] = str(sum(bool((m := lenient.search(r['generation'].strip())) and m.group(1) == r['gold'])
                                       for r in rows))
    if (V4 / 'triagem_dev' / 'manifest.json').exists():
        gate = read_json(V4 / 'triagem_dev' / 'manifest.json')
        for mode, levels in gate['summary'].items():
            for level, entry in levels.items():
                out[f'v4.gate.{mode}.{level}'] = pct(entry['accuracy'])
                out[f'v4.gate.{mode}.{level}.en'] = f"{100 * entry['accuracy']:.1f}%"
                out[f'v4.gate.{mode}.{level}.k'] = frac(entry['correct'], entry['n'])
                low, high = wilson(entry['correct'], entry['n'])
                out[f'v4.gate.{mode}.{level}.ic'] = f'{pct(low)} a {pct(high)}'
                out[f'v4.gate.{mode}.{level}.formato'] = frac(entry['format_valid'], entry['n'])
                out[f'v4.gate.{mode}.{level}.trunc'] = str(entry['truncated'])
    if VALIDATION.exists():
        for key, value in read_json(VALIDATION).get('numeros', {}).items():
            out['val.' + key] = str(value)
    return out


# ------------------------------------------------------------------ tabelas
def table_md(header, rows, align):
    """Tabela pipe com larguras relativas proporcionais ao conteúdo de cada coluna.

    Cada coluna recebe ao menos a largura da sua palavra mais longa (cerca de 70
    caracteres por largura de tabela), para que nenhuma palavra seja partida.
    """
    cells = [[str(r[i]) for r in [header] + rows] for i in range(len(header))]
    lengths = [max(9, max(len(c) for c in col)) for col in cells]
    longest = [max(len(w) for c in col for w in (c.split() or [''])) for col in cells]
    total = sum(lengths)
    dashes = [max(3, round(100 * n / total), -(-100 * (w + 2) // 70)) for n, w in zip(lengths, longest)]
    sep = {'l': lambda n: ':' + '-' * n, 'r': lambda n: '-' * n + ':', 'c': lambda n: ':' + '-' * n + ':'}
    lines = ['| ' + ' | '.join(header) + ' |', '|' + '|'.join(sep[a](n) for a, n in zip(align, dashes)) + '|']
    def cell(value):
        # Marcadores como <answer> viram código; o pandoc os descartaria como HTML.
        return re.sub(r'(<[^<>\s]+>)', r'`\1`', str(value)).replace('|', '\\|')
    for row in rows:
        lines.append('| ' + ' | '.join(cell(c) for c in row) + ' |')
    return '\n'.join(lines)


def tabelas():
    import tabelas_artigo
    return tabelas_artigo.build(ROOT, REPO, num, pct, frac)


FIGURAS = {
    'piloto': ('fig1_piloto_acuracia.png', 'Acurácia registrada no estudo-piloto por nível e alvo de treinamento',
               'Reanálise dos 540 registros publicados. Barras: média das três sementes; círculos: sementes.'),
    'triagem': ('fig2_triagem_v3.png', 'Triagem DEV do estudo v3.2: valores documentados e reproduções',
                'README do v3 (commit 273d131) e reproduções desta revisão. Linha tracejada: critério de 0,65.'),
    'falhas': ('fig3_falhas_v3.png', 'Composição das respostas da reprodução T4 da triagem DEV v3.2',
               'Classificação descritiva das 480 respostas; o desfecho registrado usa apenas o corretor congelado.'),
    'v4': ('fig4_triagem_v4.png', 'Triagem DEV do protocolo v4 antes do registro',
           'Arquivos de experimento_cientifico_v4/triagem_dev. Linha tracejada: critério de 0,65.'),
}


# ---------------------------------------------------------------- expansão
def expand(text, data, tables):
    order = {'tab': [], 'fig': []}
    for kind, key in re.findall(r'\{\{(tabela|figura):([\w-]+)\}\}', text):
        order['tab' if kind == 'tabela' else 'fig'].append(key)
    number = {('tab', k): i + 1 for i, k in enumerate(order['tab'])}
    number.update({('fig', k): i + 1 for i, k in enumerate(order['fig'])})

    def ref(match):
        kind, key = match.group(1), match.group(2)
        if (kind, key) not in number:
            raise KeyError(f'Referência sem destino: @{kind}:{key}')
        return ('Tabela ' if kind == 'tab' else 'Figura ') + str(number[kind, key])
    text = re.sub(r'@(tab|fig):([\w-]+)', ref, text)

    draft = os.environ.get('COT_RASCUNHO') == '1'

    def value(match):
        key = match.group(1)
        if key not in data:
            if draft:
                return '[' + key + ']'
            raise KeyError('Valor sem fonte de dados: ' + key)
        return data[key]
    text = re.sub(r'\{\{v:([\w.\-]+)\}\}', value, text)

    def block(match):
        kind, key = match.group(1), match.group(2)
        if kind == 'tabela':
            if key not in tables and draft:
                return f'**[tabela {key} pendente]**\n'
            title, header, rows, align, source = tables[key]
            return ('::: {custom-style="Legenda"}\nTabela %d – %s\n:::\n\n%s\n\n'
                    '::: {custom-style="Fonte"}\nFonte: %s\n:::\n' % (number['tab', key], title,
                                                                    table_md(header, rows, align), source))
        image, title, source = FIGURAS[key]
        path = ROOT / 'figures' / image
        if not path.exists():
            if draft:
                return f'**[figura {key} pendente]**\n'
            raise FileNotFoundError(path)
        return ('::: {custom-style="Legenda"}\nFigura %d – %s\n:::\n\n'
                '::: {custom-style="Figura"}\n![](%s){width=15cm}\n:::\n\n'
                '::: {custom-style="Fonte"}\nFonte: %s\n:::\n' % (number['fig', key], title, path, source))
    text = re.sub(r'^\{\{(tabela|figura):([\w-]+)\}\}\s*$', block, text, flags=re.M)
    text = text.replace('{{quebra}}', '```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```')
    leftover = [m for m in re.findall(r'\{\{[^}]*\}\}|<<[A-Z_]+>>', text) if m != '{{sumario}}']
    if leftover and not draft:
        raise ValueError('Marcadores não resolvidos: ' + ', '.join(leftover[:5]))
    return text, number


# ------------------------------------------------------- pós-processamento
def style_tables(docx_path):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    from modelo_docx import find_style
    doc = Document(str(docx_path))
    table_style = find_style(doc, 'Texto de tabela')

    def border(parent, edge, size):
        element = OxmlElement(f'w:{edge}')
        element.set(qn('w:val'), 'single' if size else 'nil')
        element.set(qn('w:sz'), str(size))
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), '000000')
        parent.append(element)

    for table in doc.tables:
        tbl_pr = table._tbl.tblPr
        for old in tbl_pr.findall(qn('w:tblBorders')) + tbl_pr.findall(qn('w:tblW')):
            tbl_pr.remove(old)
        width = OxmlElement('w:tblW')
        width.set(qn('w:w'), '5000')
        width.set(qn('w:type'), 'pct')
        tbl_pr.append(width)
        borders = OxmlElement('w:tblBorders')
        for edge, size in (('top', 10), ('left', 0), ('bottom', 10), ('right', 0),
                           ('insideH', 0), ('insideV', 0)):
            border(borders, edge, size)
        tbl_pr.append(borders)
        for index, row in enumerate(table.rows):
            tr_pr = row._tr.get_or_add_trPr()
            tr_pr.append(OxmlElement('w:cantSplit'))
            if index == 0:
                tr_pr.append(OxmlElement('w:tblHeader'))
            for cell in row.cells:
                tc_pr = cell._tc.get_or_add_tcPr()
                if index == 0:
                    tc_borders = OxmlElement('w:tcBorders')
                    border(tc_borders, 'bottom', 6)
                    tc_pr.append(tc_borders)
                for paragraph in cell.paragraphs:
                    paragraph.style = table_style
                    paragraph.paragraph_format.keep_with_next = index < len(table.rows) - 1
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
                        if index == 0:
                            run.font.bold = True
    for paragraph in doc.paragraphs:
        if paragraph.style.name == 'Figura':
            paragraph.paragraph_format.keep_with_next = True
    doc.core_properties.title = TITLE
    doc.core_properties.author = 'Kelvin e Silva Marques et al.'
    doc.core_properties.keywords = 'cadeia de pensamento; monitorabilidade; pré-registro; reprodutibilidade'
    doc.save(str(docx_path))


TITLE = ''


# --------------------------------------------------------------- verificação
def words(text):
    text = text.replace('­', '').replace('ﬁ', 'fi').replace('ﬂ', 'fl')
    text = text.replace('-\n', '-')
    return re.findall(r'[0-9A-Za-zÀ-ÿ]+', text.lower())


def docx_text(path):
    """Texto de parágrafos, tabelas e equações (OMML), na ordem do documento."""
    import zipfile
    from xml.etree import ElementTree as ET
    xml = zipfile.ZipFile(path).read('word/document.xml')
    root = ET.fromstring(xml)
    w = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    m = '{http://schemas.openxmlformats.org/officeDocument/2006/math}'
    parts = []
    for paragraph in root.iter(w + 'p'):
        parts.append(''.join((node.text or '') if node.tag == w + 't' else
                             ' ' if node.tag in (w + 'tab', w + 'br') else ' ' + (node.text or '') + ' '
                             for node in paragraph.iter() if node.tag in (w + 't', m + 't', w + 'tab', w + 'br')))
    return '\n'.join(parts)


def verify(docx_path, pdf_path):
    import modelo_docx
    pdf_text = subprocess.run(['pdftotext', '-layout', str(pdf_path), '-'], capture_output=True,
                              text=True, check=True).stdout
    # Cabeçalho e rodapé se repetem em cada página do PDF e não pertencem ao corpo do DOCX.
    pdf_text = pdf_text.replace(modelo_docx.HEADER_TEXT, ' ').replace(modelo_docx.FOOTER_TEXT, ' ')
    a, b = Counter(words(docx_text(docx_path))), Counter(words(pdf_text))
    page_numbers = {str(i) for i in range(1, 200)}
    missing = {w: c - b[w] for w, c in a.items() if c > b[w]}
    extra = {w: c - a[w] for w, c in b.items() if c > a[w] and w not in page_numbers}
    info = subprocess.run(['pdfinfo', str(pdf_path)], capture_output=True, text=True, check=True).stdout
    pages = int(re.search(r'Pages:\s+(\d+)', info).group(1))
    fonts = subprocess.run(['pdffonts', str(pdf_path)], capture_output=True, text=True).stdout.splitlines()[2:]
    not_embedded = [line for line in fonts if line.split() and line.split()[-5] != 'yes']
    return dict(pages=pages, docx_words=sum(a.values()), pdf_words=sum(b.values()),
                missing_in_pdf=dict(sorted(missing.items())[:40]), extra_in_pdf=dict(sorted(extra.items())[:40]),
                fonts=[line.split()[0] for line in fonts if line.split()], fonts_not_embedded=not_embedded,
                equivalent=not missing and not not_embedded)


def headings_after_toc(expanded):
    """Títulos de nível 1 e 2 que entram no sumário (os pré-textuais ficam antes do marcador)."""
    body = expanded.split('{{sumario}}', 1)[1]
    out = []
    for line in body.splitlines():
        match = re.match(r'^(#{1,2}) (.+?)\s*$', line)
        if match:
            out.append((len(match.group(1)), match.group(2)))
    return out


def toc_block(entries, pages):
    """Sumário estático (o DOCX permanece editável; os números vêm do PDF gerado)."""
    from xml.sax.saxutils import escape
    paragraphs = ['<w:p><w:pPr><w:pStyle w:val="Titulopretextual"/></w:pPr><w:r><w:t>SUMÁRIO</w:t></w:r></w:p>']
    for (level, title), page in zip(entries, pages):
        paragraphs.append(f'<w:p><w:pPr><w:pStyle w:val="Sumario{level}"/></w:pPr>'
                          f'<w:r><w:t xml:space="preserve">{escape(title)}</w:t></w:r><w:r><w:tab/></w:r>'
                          f'<w:r><w:t>{page}</w:t></w:r></w:p>')
    return '```{=openxml}\n' + '\n'.join(paragraphs) + '\n```'


def find_pages(pdf_path, entries):
    """Página física de cada título, procurada depois da página do sumário."""
    info = subprocess.run(['pdfinfo', str(pdf_path)], capture_output=True, text=True, check=True).stdout
    total = int(re.search(r'Pages:\s+(\d+)', info).group(1))
    texts = [subprocess.run(['pdftotext', '-f', str(i), '-l', str(i), '-layout', str(pdf_path), '-'],
                            capture_output=True, text=True, check=True).stdout for i in range(1, total + 1)]
    start = next(i for i, page in enumerate(texts) if re.search(r'^\s*SUMÁRIO\s*$', page, re.M)) + 1
    pages, cursor = [], start
    for _, title in entries:
        pattern = re.compile(r'^\s*' + r'\s+'.join(map(re.escape, title.split())) + r'\s*$', re.M)
        found = next((i for i in range(cursor, total) if pattern.search(texts[i])), None)
        if found is None:
            raise RuntimeError('Título não localizado no PDF para o sumário: ' + title)
        pages.append(found + 1)
        cursor = found
    return pages


def render(expanded, reference):
    """Pandoc (DOCX) e LibreOffice (PDF a partir do mesmo DOCX)."""
    (BUILD / 'artigo_expandido.md').write_text(expanded, encoding='utf-8')
    docx_path = DOCS / f'{NAME}.docx'
    subprocess.run(['pandoc', str(BUILD / 'artigo_expandido.md'), '-f', 'markdown+pipe_tables', '-t', 'docx',
                    '--reference-doc', str(reference), '-o', str(docx_path)], check=True)
    style_tables(docx_path)
    lo_profile = BUILD / 'lo-profile'
    subprocess.run(['soffice', f'-env:UserInstallation=file://{lo_profile}', '--headless', '--convert-to', 'pdf',
                    '--outdir', str(BUILD), str(docx_path)], check=True, capture_output=True, timeout=600)
    pdf_path = DOCS / f'{NAME}.pdf'
    shutil.move(str(BUILD / f'{NAME}.pdf'), pdf_path)
    return docx_path, pdf_path


def main():
    global TITLE
    import gerar_figuras
    import modelo_docx
    gerar_figuras.main()
    BUILD.mkdir(exist_ok=True)
    DOCS.mkdir(exist_ok=True)
    text = SOURCE.read_text(encoding='utf-8')
    title = re.search(r'^::: \{custom-style="Capa titulo"\}\n(.+?)\n:::', text, re.M | re.S).group(1).strip()
    subtitle = re.search(r'^::: \{custom-style="Capa subtitulo"\}\n(.+?)\n:::', text, re.M | re.S)
    TITLE = title + (': ' + subtitle.group(1).strip() if subtitle else '')
    expanded, numbers = expand(text, values(), tabelas())
    reference = modelo_docx.build(BUILD / 'modelo.docx')
    entries = headings_after_toc(expanded)
    pages = ['00'] * len(entries)
    for _ in range(3):  # o sumário tem tamanho fixo; a segunda passagem já é estável
        docx_path, pdf_path = render(expanded.replace('{{sumario}}', toc_block(entries, pages)), reference)
        found = find_pages(pdf_path, entries)
        if found == pages:
            break
        pages = found
    else:
        sys.exit('Números de página do sumário não estabilizaram')
    check = verify(docx_path, pdf_path)
    check['sumario'] = [dict(titulo=title, pagina=page) for (_, title), page in zip(entries, pages)]
    pages = BUILD / 'paginas'
    shutil.rmtree(pages, ignore_errors=True)
    pages.mkdir()
    subprocess.run(['pdftoppm', '-r', '70', '-png', str(pdf_path), str(pages / 'p')], check=True)
    check.update(docx_sha256=hashlib.sha256(docx_path.read_bytes()).hexdigest(),
                 pdf_sha256=hashlib.sha256(pdf_path.read_bytes()).hexdigest(),
                 source_sha256=hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
                 tables=len([k for k in numbers if k[0] == 'tab']), figures=len([k for k in numbers if k[0] == 'fig']))
    (DOCS / 'verificacao_documentos.json').write_text(json.dumps(check, indent=2, ensure_ascii=False) + '\n',
                                                      encoding='utf-8')
    print(json.dumps(check, indent=2, ensure_ascii=False))
    if not check['equivalent']:
        sys.exit('DOCX e PDF divergem ou há fontes não incorporadas; veja verificacao_documentos.json')


if __name__ == '__main__':
    main()
