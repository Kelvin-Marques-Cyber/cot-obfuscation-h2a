"""Cria o modelo de estilos (reference.docx) usado pelo pandoc para o artigo.

Padrão: A4; margens de 3 cm (superior e esquerda) e 2 cm (inferior e direita),
conforme a ABNT NBR 14724; Liberation Serif 11 pt (métrica equivalente a Times
New Roman), entrelinha 1,3; títulos em preto, numeração de página no rodapé.
Somente preto, branco e tons de cinza.
"""
from __future__ import annotations

import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = 'Liberation Serif'
MONO = 'Liberation Mono'
BLACK = RGBColor(0, 0, 0)


def _font(style, size, bold=None, italic=None, name=FONT):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = BLACK
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.append(fonts)
    for attribute in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        fonts.set(qn(attribute), name)
    for attribute in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
        fonts.attrib.pop(qn(attribute), None)


def _paragraph(style, before=0, after=6, line=1.3, align=None, indent=None, keep_next=False):
    fmt = style.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.line_spacing = line
    fmt.widow_control = True
    fmt.keep_with_next = keep_next
    if align is not None:
        fmt.alignment = align
    if indent is not None:
        fmt.first_line_indent = indent


def find_style(doc, name):
    """Procura pelo nome gravado no XML (o pandoc grava "Heading 1" com maiúscula)."""
    for style in doc.styles:
        if style.name == name:
            return style
    raise KeyError(name)


def get_style(doc, name, kind=WD_STYLE_TYPE.PARAGRAPH, base='Normal'):
    try:
        return find_style(doc, name)
    except KeyError:
        style = doc.styles.add_style(name, kind)
        if kind == WD_STYLE_TYPE.PARAGRAPH and base:
            style.base_style = find_style(doc, base)
        return style


HEADER_TEXT = 'OFUSCAÇÃO DA CoT SOB SUPERVISÃO | REANÁLISE E TRIAGEM DE COMPETÊNCIA'
FOOTER_TEXT = 'Marques et al. | 2026'


def running_header(section):
    """Cabeçalho corrente, como no artigo original; a capa (primeira página) não o recebe."""
    header = section.header
    paragraph = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(HEADER_TEXT)
    run.font.name = FONT
    run.font.size = Pt(9)


def page_number_footer(section):
    from docx.enum.text import WD_TAB_ALIGNMENT
    footer = section.footer
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT)
    label = paragraph.add_run(FOOTER_TEXT + '\t')
    label.font.name = FONT
    label.font.size = Pt(9)
    run = paragraph.add_run()
    for kind, text in (('begin', None), (None, 'PAGE'), ('end', None)):
        if kind:
            element = OxmlElement('w:fldChar')
            element.set(qn('w:fldCharType'), kind)
        else:
            element = OxmlElement('w:instrText')
            element.set(qn('xml:space'), 'preserve')
            element.text = text
        run._r.append(element)
    run.font.name = FONT
    run.font.size = Pt(10)


def build(target):
    target = Path(target)
    target.parent.mkdir(parents=True, exist_ok=True)
    default = subprocess.run(['pandoc', '-o', str(target), '--print-default-data-file', 'reference.docx'],
                             capture_output=True)
    if default.returncode:
        sys.exit(default.stderr.decode())
    doc = Document(str(target))
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.top_margin, section.left_margin = Cm(3.0), Cm(3.0)
    section.bottom_margin, section.right_margin = Cm(2.0), Cm(2.0)
    section.footer_distance = Cm(1.0)
    section.header_distance = Cm(1.5)
    # Capa sem cabeçalho nem número; as demais páginas são numeradas fisicamente.
    section.different_first_page_header_footer = True
    running_header(section)
    page_number_footer(section)

    styles = doc.styles
    normal = find_style(doc, 'Normal')
    _font(normal, 11)
    _paragraph(normal, after=0, line=1.3)
    for name in ('Body Text', 'First Paragraph'):
        style = get_style(doc, name)
        _font(style, 11)
        _paragraph(style, after=6, line=1.3, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=Cm(1.25))
    compact = get_style(doc, 'Compact')
    _font(compact, 11)
    _paragraph(compact, after=2, line=1.2, align=WD_ALIGN_PARAGRAPH.LEFT)
    sizes = {1: 12, 2: 11.5, 3: 11}
    for level, size in sizes.items():
        style = find_style(doc, f'Heading {level}')
        _font(style, size, bold=True, italic=False)
        _paragraph(style, before=14 if level == 1 else 10, after=6, line=1.15,
                   align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=True)
        style.paragraph_format.first_line_indent = Cm(0)
    for name in ('Title', 'Subtitle', 'Author', 'Date', 'Abstract'):
        style = get_style(doc, name)
        _font(style, 11)
        style.font.color.rgb = BLACK
    table_text = get_style(doc, 'Texto de tabela')
    _font(table_text, 9)
    _paragraph(table_text, after=0, line=1.05, align=WD_ALIGN_PARAGRAPH.LEFT, indent=Cm(0))
    caption = get_style(doc, 'Legenda')
    _font(caption, 10, bold=False)
    _paragraph(caption, before=10, after=4, line=1.15, align=WD_ALIGN_PARAGRAPH.LEFT,
               indent=Cm(0), keep_next=True)
    source = get_style(doc, 'Fonte')
    _font(source, 9)
    _paragraph(source, before=3, after=10, line=1.1, align=WD_ALIGN_PARAGRAPH.LEFT, indent=Cm(0))
    figure = get_style(doc, 'Figura')
    _paragraph(figure, before=2, after=2, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER,
               indent=Cm(0), keep_next=True)
    reference = get_style(doc, 'Referencia')
    _font(reference, 11)
    _paragraph(reference, after=9, line=1.0, align=WD_ALIGN_PARAGRAPH.LEFT, indent=Cm(0))
    cover = get_style(doc, 'Capa')
    _font(cover, 11)
    _paragraph(cover, after=0, line=1.3, align=WD_ALIGN_PARAGRAPH.CENTER, indent=Cm(0))
    cover_title = get_style(doc, 'Capa titulo')
    _font(cover_title, 15, bold=True)
    _paragraph(cover_title, before=120, after=8, line=1.2, align=WD_ALIGN_PARAGRAPH.CENTER, indent=Cm(0))
    cover_subtitle = get_style(doc, 'Capa subtitulo')
    _font(cover_subtitle, 12, bold=False)
    _paragraph(cover_subtitle, before=0, after=8, line=1.2, align=WD_ALIGN_PARAGRAPH.CENTER, indent=Cm(0))
    cover_place = get_style(doc, 'Capa local')
    _font(cover_place, 11, bold=False)
    _paragraph(cover_place, before=190, after=0, line=1.3, align=WD_ALIGN_PARAGRAPH.CENTER, indent=Cm(0))
    pre_title = get_style(doc, 'Titulo pretextual')
    _font(pre_title, 12, bold=True)
    _paragraph(pre_title, before=0, after=12, line=1.15, align=WD_ALIGN_PARAGRAPH.CENTER, indent=Cm(0))
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
    for level, bold in ((1, True), (2, False)):
        toc = get_style(doc, f'Sumario {level}')
        _font(toc, 11, bold=bold)
        _paragraph(toc, before=6 if level == 1 else 0, after=2, line=1.2, align=WD_ALIGN_PARAGRAPH.LEFT,
                   indent=Cm(0))
        toc.paragraph_format.left_indent = Cm(0 if level == 1 else 0.6)
        toc.paragraph_format.tab_stops.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    resumo = get_style(doc, 'Resumo')
    _font(resumo, 11)
    _paragraph(resumo, after=6, line=1.0, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=Cm(0))
    code = get_style(doc, 'Source Code')
    _font(code, 9, name=MONO)
    for name in ('Verbatim Char',):
        try:
            _font(find_style(doc, name), 9, name=MONO)
        except KeyError:
            pass
    hyperlink = None
    try:
        hyperlink = find_style(doc, 'Hyperlink')
    except KeyError:
        pass
    if hyperlink is not None:
        hyperlink.font.color.rgb = BLACK
        hyperlink.font.underline = False
    for name in ('Table Caption', 'Image Caption', 'Caption'):
        try:
            style = find_style(doc, name)
            _font(style, 10, italic=False)
        except KeyError:
            pass
    doc.core_properties.author = 'Kelvin e Silva Marques'
    doc.core_properties.comments = ''
    doc.save(str(target))
    return target


if __name__ == '__main__':
    print(build(sys.argv[1] if len(sys.argv) > 1 else 'build/modelo.docx'))
